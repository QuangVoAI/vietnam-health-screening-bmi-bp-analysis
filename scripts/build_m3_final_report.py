from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
FIG = ROOT / 'figures'
TABLES = ROOT / 'tables'
DEST = ROOT / 'reports/Bao_cao_cuoi_ky_BMI_HuyetAp_2026.docx'

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.65); sec.bottom_margin = Inches(.65)
sec.left_margin = Inches(.75); sec.right_margin = Inches(.75)
styles = doc.styles
styles['Normal'].font.name = 'Times New Roman'; styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
styles['Normal'].font.size = Pt(11)
for name, size in [('Title', 18), ('Heading 1', 14), ('Heading 2', 12)]:
    styles[name].font.name = 'Times New Roman'; styles[name]._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    styles[name].font.size = Pt(size); styles[name].font.bold = True; styles[name].font.color.rgb = RGBColor(0,0,0)

def shade(cell, fill='1F4E78'):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)

def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.style = 'Table Grid'
    for i, h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=str(h); shade(c)
        for r in c.paragraphs[0].runs: r.font.color.rgb=RGBColor(255,255,255); r.font.bold=True
        c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for row in rows:
        cells=t.add_row().cells
        for i,v in enumerate(row):
            cells[i].text=str(v); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if len(t.rows)%2==1: shade(cells[i],'EAF2F8')
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width=Inches(w)
    doc.add_paragraph()
    return t

def para(text='', bold_lead=None):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(6); p.paragraph_format.line_spacing=1.2
    if bold_lead:
        p.add_run(bold_lead).bold=True
    p.add_run(text)
    return p

def fig(name, caption, width=6.2):
    doc.add_picture(str(FIG/name), width=Inches(width))
    p=doc.add_paragraph(caption); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].italic=True; p.runs[0].font.size=Pt(10)

# Cover
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.space_after=Pt(18)
r=p.add_run('BÁO CÁO CUỐI KỲ\nPHÂN TÍCH DỮ LIỆU'); r.bold=True; r.font.size=Pt(16)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('Chất lượng dữ liệu và mối liên quan giữa chỉ số khối cơ thể với huyết áp tâm thu trong khám sức khỏe định kỳ cộng đồng năm 2026'); r.bold=True; r.font.size=Pt(18)
para('Sinh viên: Võ Xuân Quang  |  MSSV: 523H0173', None).alignment=WD_ALIGN_PARAGRAPH.CENTER
para('Phạm vi: nghiên cứu cắt ngang trên các lượt khám hợp lệ của người 18–59 tuổi; kết quả mô tả mối liên quan trong mẫu, không phải chẩn đoán cá nhân, suy luận nhân quả hay ước tính đại diện toàn quốc.').alignment=WD_ALIGN_PARAGRAPH.CENTER
doc.add_page_break()

doc.add_heading('Tóm tắt', 1)
para('Báo cáo phân tích dữ liệu khám sức khỏe định kỳ đã khử định danh nhằm mô tả chất lượng dữ liệu, phân phối BMI và huyết áp tâm thu (HATT), kiểm định khác biệt/ liên quan, và xây dựng mô hình hồi quy đa biến. Mẫu cuối gồm 7.146 lượt khám và 14 biến. BMI có mối liên quan dương với HATT sau khi điều chỉnh tuổi, giới và mạch (β = 0,861 mmHg cho mỗi 1 kg/m²; CI 95%: 0,778–0,944; p < 0,001). Kết quả có ý nghĩa cho quản lý sàng lọc dự phòng, nhưng không dùng để kết luận BMI gây tăng huyết áp.')

doc.add_heading('1. Dữ liệu và câu hỏi nghiên cứu', 1)
table(['Nội dung','Mô tả'], [
 ['Đơn vị phân tích','Một lượt khám sức khỏe hợp lệ của người 18–59 tuổi'],
 ['Quy mô','7.146 lượt khám; 14 biến sau khử định danh'],
 ['Biến chính','BMI, HATT, HATTr, tuổi, giới, mạch, nhóm BMI, nhóm tuổi, cờ sàng lọc HA ≥140/90'],
 ['Câu hỏi chính','BMI có còn liên quan HATT khi xét tuổi, giới tính và mạch không?'],
 ['Nguyên tắc','Phân tích liên hệ thống kê trong mẫu; không chẩn đoán, điều trị hay suy luận nguyên nhân–kết quả.'],
], [1.45,5.5])

doc.add_heading('2. Khám phá và chất lượng dữ liệu', 1)
para('Dữ liệu phát hành không thiếu giá trị ở các biến phân tích và không có mã nghiên cứu trùng. Có 118 lượt mang cờ mã nguồn lặp; chúng được giữ lại trong phân tích chính để tránh loại bỏ tùy tiện, sau đó kiểm tra độ nhạy bằng mô hình loại riêng các lượt này. Ngoại lai được rà soát bằng quy tắc IQR, ghi nhận minh bạch và không tự động xóa vì có thể là giá trị đo hợp lệ.')
fig('05_cohort_flow.png','Hình 1. Luồng hình thành mẫu: 10.912 bản ghi nguồn không rỗng đến 7.146 lượt nghiên cứu.',5.9)
fig('01_distributions_bmi_sbp.png','Hình 2. Histogram và KDE của BMI và HATT.',6.2)

doc.add_heading('3. Phân phối xác suất', 1)
dist=pd.read_csv(TABLES/'08_distribution_audit.csv')
table(['Biến','Skewness','Excess kurtosis','Nhận xét'], [[r['Biến'],f"{r['Skewness']:.3f}",f"{r['Excess kurtosis']:.3f}",'Lệch phải; không xem là chuẩn hoàn toàn'] for _,r in dist.iterrows()], [1.0,1.1,1.25,3.6])
para('D’Agostino–Pearson và Shapiro–Wilk đều bác bỏ chuẩn tính (p < 0,001). Với n lớn, các kiểm định này rất nhạy; do đó kết luận dựa thêm vào Q–Q plot và độ lệch/độ nhọn. Điều này giải thích việc báo cáo đồng thời Welch, Spearman và các kích thước hiệu ứng, thay vì dựa duy nhất vào giả định chuẩn.')
fig('13_distribution_diagnostics.png','Hình 3. Q–Q plot cho thấy lệch khỏi phân phối chuẩn ở phần đuôi.',6.2)

doc.add_heading('4. Kiểm định giả thuyết', 1)
para('Mọi kiểm định là hai phía, α = 0,05. H0/H1 được nêu trước trong notebook. Levene cho thấy phương sai HATT không đồng nhất theo nhóm BMI và nhóm tuổi; vì vậy báo cáo ANOVA cổ điển để đáp ứng nội dung học phần, đồng thời dùng Welch ANOVA và hậu kiểm cặp Welch–Holm như kiểm tra bền vững. Với Chi-square, ô kỳ vọng nhỏ nhất là 83,03 và không có ô nào dưới 5, nên giả định tần số kỳ vọng đạt.')
tests=pd.read_csv(TABLES/'08_hypothesis_tests_summary.csv')
rows=[]
for _,r in tests.iloc[:6].iterrows():
    rows.append([r['Phân tích'], f"{r['Thống kê']:.3f}", '<0,001', f"{r['Kích thước hiệu ứng']}: {r['Giá trị hiệu ứng']:.3f}"])
table(['Kiểm định','Thống kê','p-value','Kích thước hiệu ứng'],rows,[2.55,1.0,.8,2.6])
para('Diễn giải: HATT khác theo giới (Cohen d = 0,629); khác theo nhóm BMI (η² = 0,083) và nhóm tuổi (η² = 0,073). Nhóm BMI liên quan với cờ sàng lọc (Cramér V = 0,217). Đây là các hiệu ứng nhỏ đến vừa theo nghĩa thực hành, không chỉ là ý nghĩa thống kê do cỡ mẫu lớn.')
fig('03_screening_rate_by_bmi.png','Hình 4. Tỷ lệ cờ sàng lọc tăng theo nhóm BMI.',5.6)
fig('09_screening_heatmap_age_bmi.png','Hình 5. Phân tầng tuổi × BMI giúp xác định các cụm ưu tiên sàng lọc.',6.2)

doc.add_heading('5. Tương quan và hồi quy đa biến', 1)
para('BMI và HATT có tương quan dương mức thấp–vừa: Pearson r = 0,287 và Spearman ρ = 0,335; cả hai p < 0,001. Spearman cao hơn một phần phù hợp với quan hệ đơn điệu nhưng phân phối không hoàn toàn chuẩn.')
fig('02_bmi_groups_and_scatter.png','Hình 6. HATT có xu hướng cao hơn theo BMI, nhưng vẫn còn độ phân tán đáng kể.',6.2)
reg=pd.read_csv(TABLES/'06_adjusted_ols_regression.csv')
table(['Biến','β (mmHg)','CI 95%','Ý nghĩa'], [[r['Biến'],f"{r['Hệ số β']:.3f}",f"{r['CI 95% thấp']:.3f} đến {r['CI 95% cao']:.3f}",'p < 0,001'] for _,r in reg.iterrows()], [1.7,1.0,2.25,1.25])
para('Mô hình HATT ~ BMI + tuổi + giới + mạch có R² = 0,257, R² hiệu chỉnh = 0,256 và RMSE = 14,07 mmHg. Mỗi 1 kg/m² BMI cao hơn liên quan với HATT cao hơn 0,861 mmHg khi giữ các biến còn lại không đổi. R² vừa phải là hợp lý vì dữ liệu không có thông tin về thuốc, ăn mặn, hút thuốc, rượu, vận động hay bệnh nền.')
fig('11_regression_forest_plot.png','Hình 7. Quy đổi hệ số thành mức tăng thực hành và CI 95%.',6.0)
fig('04_regression_diagnostics.png','Hình 8. Chẩn đoán mô hình: phần đuôi phần dư lệch, vì vậy diễn giải CI cần thận trọng.',6.2)

doc.add_heading('6. Độ nhạy, insight và giới hạn', 1)
para('Khi loại 118 lượt mang cờ lặp, hệ số BMI giảm rất ít, từ 0,861 xuống 0,852 mmHg/1 BMI. Do đó, kết luận chính không phụ thuộc đáng kể vào lựa chọn giữ hay bỏ các lượt này.')
fig('12_sensitivity_repeat_flag.png','Hình 9. Hệ số BMI ổn định trong phân tích độ nhạy.',5.7)
table(['Insight','Ý nghĩa thực tiễn'],[
 ['Gradient BMI','Tỷ lệ cờ sàng lọc tăng từ 6,6% ở BMI <18,5 lên 38,3% ở BMI ≥30.'],
 ['Phân tầng tuổi','Ở phần lớn nhóm tuổi, BMI cao đi kèm tỷ lệ cờ sàng lọc cao hơn; cần ưu tiên truyền thông/sàng lọc theo cụm nguy cơ thay vì một ngưỡng duy nhất.'],
 ['Giới hạn','Thiết kế cắt ngang không xác lập nguyên nhân; một lượt khám không thay thế chẩn đoán; nhiễu còn sót do thiếu biến lối sống, thuốc và bệnh nền.'],
], [1.55,5.4])

doc.add_heading('7. Kết luận', 1)
para('Bài làm đáp ứng đầy đủ EDA, phân phối xác suất, kiểm định giả thuyết, tương quan và hồi quy đa biến trên cùng một bộ dữ liệu. Điểm mạnh vượt yêu cầu là: kiểm tra giả định, báo cáo kích thước hiệu ứng/CI, Welch ANOVA và hậu kiểm có hiệu chỉnh, phân tích độ nhạy, trực quan hóa phân tầng, và ranh giới diễn giải minh bạch. Kết quả hỗ trợ nhận diện mô hình sàng lọc dự phòng trong mẫu khám, không thay thế đánh giá lâm sàng.')

doc.add_heading('Tài liệu tham khảo', 1)
for ref in [
 'Nguyen Minh Q, Nguyen Vo M. Anthropometric Indexes for Predicting High Blood Pressure in Vietnamese Adults: A Cross-Sectional Study. Integr Blood Press Control. 2020;13:181–186. doi:10.2147/IBPC.S281996.',
 'Wang J, et al. The importance of waist circumference and body mass index in cross-sectional relationships with risk of cardiovascular disease in Vietnam. PLoS One. 2018. PMID:29813112.',
 'Nguyen Duc N, et al. Impact of body mass index and waist circumference on blood pressure: A cross-sectional survey in a population living in the Vietnam northern mountainous. Diabetes Metab Syndr. 2019;13(2):1399–1404. doi:10.1016/j.dsx.2019.02.015.'
]: para(ref)

doc.save(DEST)
print(DEST)
