"""Create a template-derived academic report for the BMI and blood-pressure study."""
from pathlib import Path
import os
import shutil
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = Path(os.environ['M3_REPORT_TEMPLATE'])
FIG = ROOT / 'figures'; TAB = ROOT / 'tables'
DEST = ROOT / 'reports/Research_Report_BMI_Blood_Pressure_2026_EN.docx'
WORK = ROOT / '.report_build/Research_Report_BMI_Blood_Pressure_2026_EN_working.docx'

if not TEMPLATE.exists():
    raise FileNotFoundError('Set M3_REPORT_TEMPLATE to an accessible DOCX template.')
WORK.parent.mkdir(parents=True, exist_ok=True)

shutil.copy2(TEMPLATE, WORK)
doc = Document(WORK)

# Clear only body contents; preserve template package, styles, theme and section properties.
body = doc._element.body
for child in list(body):
    if child.tag != qn('w:sectPr'):
        body.remove(child)

for s in doc.sections:
    s.page_width = Inches(8.27); s.page_height = Inches(11.69)
    s.left_margin = Inches(1.38); s.right_margin = Inches(.79)
    s.top_margin = Inches(1.38); s.bottom_margin = Inches(1.18)

styles = doc.styles
for name in ['Normal','Body Text','Heading 1','Heading 2','Caption']:
    st = styles[name]
    st.font.name = 'Times New Roman'; st._element.rPr.rFonts.set(qn('w:eastAsia'),'Times New Roman')
styles['Normal'].font.size = Pt(12)
styles['Body Text'].font.size = Pt(12)
styles['Heading 1'].font.size = Pt(14); styles['Heading 1'].font.bold = True
styles['Heading 2'].font.size = Pt(12); styles['Heading 2'].font.bold = True
styles['Caption'].font.size = Pt(10); styles['Caption'].font.italic = True
for name, size in [('Research Heading 1', 14), ('Research Heading 2', 12)]:
    if name not in [s.name for s in styles]:
        st = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        st.base_style = styles['Normal']
    st = styles[name]
    st.font.name = 'Times New Roman'; st._element.rPr.rFonts.set(qn('w:eastAsia'),'Times New Roman')
    st.font.size = Pt(size); st.font.bold = True

def p(text='', style='Normal', align=None, before=0, after=6, bold=False, italic=False):
    x = doc.add_paragraph(style=style); x.paragraph_format.space_before=Pt(before); x.paragraph_format.space_after=Pt(after); x.paragraph_format.line_spacing=1.25
    if align is not None: x.alignment=align
    r=x.add_run(text); r.bold=bold; r.italic=italic
    return x

def heading(text, level=1):
    return doc.add_paragraph(text, style='Research Heading 1' if level == 1 else 'Research Heading 2')

def page_break(): doc.add_page_break()

def shade(cell, fill='D9EAF7'):
    tcPr=cell._tc.get_or_add_tcPr(); e=OxmlElement('w:shd'); e.set(qn('w:fill'),fill); tcPr.append(e)

def table(headers, rows, size=9.5):
    t=doc.add_table(rows=1, cols=len(headers)); t.style='Table Grid'; t.autofit=True
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=str(h); shade(c,'1F4E78')
        for run in c.paragraphs[0].runs: run.font.bold=True; run.font.color.rgb=None; run.font.size=Pt(size)
    trPr = t.rows[0]._tr.get_or_add_trPr(); tblHeader = OxmlElement('w:tblHeader'); tblHeader.set(qn('w:val'), 'true'); trPr.append(tblHeader)
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        rowPr = t.rows[-1]._tr.get_or_add_trPr(); cantSplit = OxmlElement('w:cantSplit'); rowPr.append(cantSplit)
        for i,val in enumerate(row):
            cells[i].text=str(val)
            for par in cells[i].paragraphs:
                for run in par.runs: run.font.size=Pt(size)
            if ri%2==1: shade(cells[i],'EDF4FA')
    p('',after=3)
    return t

def figure(filename, caption, width=5.8):
    path=FIG/filename
    if path.exists():
        x=doc.add_paragraph(); x.alignment=WD_ALIGN_PARAGRAPH.CENTER
        x.add_run().add_picture(str(path), width=Inches(width))
        p(caption,style='Caption',align=WD_ALIGN_PARAGRAPH.CENTER,after=10)

def equation(label, expression):
    # Native Word math paragraph: linear Cambria Math representation is editable in Word.
    para=doc.add_paragraph(); para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    omathpara=OxmlElement('m:oMathPara'); omath=OxmlElement('m:oMath'); run=OxmlElement('m:r'); txt=OxmlElement('m:t'); txt.text=expression
    run.append(txt); omath.append(run); omathpara.append(omath); para._p.append(omathpara)
    p(label,style='Caption',align=WD_ALIGN_PARAGRAPH.CENTER,after=8)

def bullets(items):
    for item in items:
        x=doc.add_paragraph(style='Normal'); x.paragraph_format.left_indent=Inches(.25); x.paragraph_format.first_line_indent=Inches(-.16); x.paragraph_format.space_after=Pt(2); x.add_run('• ' + item)

# Cover 1
p('VIETNAM GENERAL CONFEDERATION OF LABOR',align=WD_ALIGN_PARAGRAPH.CENTER,bold=True,after=1)
p('TON DUC THANG UNIVERSITY',align=WD_ALIGN_PARAGRAPH.CENTER,bold=True,after=1)
p('FACULTY OF INFORMATION TECHNOLOGY',align=WD_ALIGN_PARAGRAPH.CENTER,bold=True,after=35)
p('VÕ XUÂN QUANG – 523H0173',align=WD_ALIGN_PARAGRAPH.CENTER,bold=True,after=55)
p('FINAL COURSE REPORT',align=WD_ALIGN_PARAGRAPH.CENTER,bold=True,after=12)
p('DATA ANALYSIS',align=WD_ALIGN_PARAGRAPH.CENTER,bold=True,after=26)
p('DATA QUALITY AND THE ASSOCIATION BETWEEN BODY MASS INDEX AND SYSTOLIC BLOOD PRESSURE IN COMMUNITY PERIODIC HEALTH SCREENING IN 2026',align=WD_ALIGN_PARAGRAPH.CENTER,bold=True,after=55)
p('HO CHI MINH CITY, 2026',align=WD_ALIGN_PARAGRAPH.CENTER,bold=True)
page_break()

# Cover 2
p('VIETNAM GENERAL CONFEDERATION OF LABOR',align=WD_ALIGN_PARAGRAPH.CENTER,bold=True,after=1)
p('TON DUC THANG UNIVERSITY',align=WD_ALIGN_PARAGRAPH.CENTER,bold=True,after=1)
p('FACULTY OF INFORMATION TECHNOLOGY',align=WD_ALIGN_PARAGRAPH.CENTER,bold=True,after=32)
p('VÕ XUÂN QUANG – 523H0173',align=WD_ALIGN_PARAGRAPH.CENTER,bold=True,after=42)
p('FINAL COURSE REPORT',align=WD_ALIGN_PARAGRAPH.CENTER,bold=True,after=10)
p('DATA ANALYSIS',align=WD_ALIGN_PARAGRAPH.CENTER,bold=True,after=22)
p('DATA QUALITY AND THE ASSOCIATION BETWEEN BODY MASS INDEX AND SYSTOLIC BLOOD PRESSURE IN COMMUNITY PERIODIC HEALTH SCREENING IN 2026',align=WD_ALIGN_PARAGRAPH.CENTER,bold=True,after=40)
p('Advised by',align=WD_ALIGN_PARAGRAPH.CENTER,after=3)
p('Course instructor',align=WD_ALIGN_PARAGRAPH.CENTER,bold=True,after=45)
p('HO CHI MINH CITY, 2026',align=WD_ALIGN_PARAGRAPH.CENTER,bold=True)
page_break()

heading('ACKNOWLEDGEMENT',1)
p('I respectfully acknowledge the guidance provided by the course instructor and the Faculty of Information Technology. This report was prepared for academic assessment in Data Analysis. The study uses de-identified screening data and is presented solely for educational and research purposes.')
p('Ho Chi Minh City, October 2026',align=WD_ALIGN_PARAGRAPH.RIGHT,after=2)
p('Author',align=WD_ALIGN_PARAGRAPH.RIGHT,after=22)
p('Võ Xuân Quang',align=WD_ALIGN_PARAGRAPH.RIGHT,bold=True)
page_break()

heading('DECLARATION OF AUTHORSHIP',1)
p('I declare that this report is my own academic work. The analysis, interpretation, figures, tables, and conclusions were prepared for this course using the stated analytical procedures. Published studies are cited where their ideas or findings are discussed. The released research materials exclude direct identifiers and do not contain the original identifiable source records.')
p('I understand that the results describe statistical associations in the study sample. They are not individual diagnoses, treatment recommendations, causal conclusions, or nationally representative estimates.')
p('Ho Chi Minh City, October 2026',align=WD_ALIGN_PARAGRAPH.RIGHT,after=2)
p('Author',align=WD_ALIGN_PARAGRAPH.RIGHT,after=22)
p('Võ Xuân Quang',align=WD_ALIGN_PARAGRAPH.RIGHT,bold=True)
page_break()

heading('ABSTRACT',1)
p('Background: Body mass index (BMI) and blood pressure are routinely collected in community health screening, yet their analytical value depends on data quality, transparent statistical assumptions, and cautious interpretation. Objective: To describe the quality and distribution of de-identified periodic screening data and to evaluate the association between BMI and systolic blood pressure (SBP) after accounting for age, sex, and pulse rate. Methods: This cross-sectional secondary analysis included 7,146 valid screening encounters among adults aged 18–59 years. Exploratory analysis, distribution diagnostics, Welch t-tests, ANOVA and Welch ANOVA, Chi-square tests, Pearson and Spearman correlations, multiple linear regression, residual diagnostics, and a repeat-record sensitivity analysis were performed. Results: BMI and SBP were right-skewed and not completely normally distributed. BMI was positively correlated with SBP (Pearson r = 0.287; Spearman ρ = 0.335; both p < 0.001). In the adjusted model, each 1 kg/m² higher BMI was associated with 0.861 mmHg higher SBP (95% CI: 0.778–0.944; p < 0.001). The screening-flag proportion increased from 6.6% for BMI <18.5 to 38.3% for BMI ≥30. Excluding 118 source-repeat-flagged encounters produced a similar BMI coefficient (0.852 mmHg). Conclusion: BMI was statistically associated with SBP and the screening flag within this sample. Findings support screening prioritisation, not diagnosis or causal inference.')
p('Keywords: body mass index; systolic blood pressure; community screening; cross-sectional analysis; multiple linear regression.',italic=True)
page_break()

heading('TABLE OF CONTENTS',1)
for line in ['ABSTRACT','LIST OF ABBREVIATIONS','CHAPTER 1. INTRODUCTION','CHAPTER 2. LITERATURE REVIEW AND CONCEPTUAL BASIS','CHAPTER 3. DATA AND METHODS','CHAPTER 4. RESULTS','CHAPTER 5. DISCUSSION','CHAPTER 6. CONCLUSION AND RECOMMENDATIONS','REFERENCES','APPENDICES']:
    p(line,after=3)
page_break()

heading('LIST OF ABBREVIATIONS',1)
table(['Abbreviation','Meaning'],[
 ['BMI','Body mass index'],['SBP','Systolic blood pressure'],['DBP','Diastolic blood pressure'],['IQR','Interquartile range'],['CI','Confidence interval'],['VIF','Variance inflation factor'],['OLS','Ordinary least squares'],['EDA','Exploratory data analysis']])
page_break()

heading('CHAPTER 1. INTRODUCTION',1)
heading('1.1 Background',2)
p('Cardiometabolic risk prevention depends partly on timely identification of modifiable risk patterns. Periodic community screening commonly records height, weight, pulse, and blood pressure, creating an opportunity to describe population-facing screening patterns. BMI is inexpensive and easy to calculate, while SBP is a clinically meaningful continuous measurement. However, neither metric alone provides a diagnosis, and a rigorous analysis must distinguish statistical association from causation or clinical decision-making.')
heading('1.2 Problem statement and research gap',2)
p('Simple descriptive counts of overweight or elevated blood pressure do not answer whether BMI remains associated with SBP after accounting for demographic and physiological covariates. Moreover, health-screening records require explicit quality checks, distribution assessment, outlier auditing, and robustness analysis before they can support a defensible conclusion.')
heading('1.3 Objectives and research questions',2)
bullets(['Describe completeness, duplicate status, outlier flags, and the distributions of key variables.', 'Test whether SBP differs by sex, BMI group, and age group.', 'Test whether BMI group is associated with the SBP/DBP screening flag.', 'Estimate Pearson and Spearman associations between BMI and SBP.', 'Estimate the adjusted association using the model SBP ~ BMI + age + sex + pulse.', 'Evaluate whether the BMI coefficient remains stable after excluding repeat-flagged encounters.'])
heading('1.4 Scope and contribution',2)
p('The unit of analysis is one valid screening encounter. The report contributes an auditable workflow combining EDA, assumption-aware inference, effect sizes and confidence intervals, stratified visualisation, and sensitivity analysis. It does not estimate causal effects, diagnose hypertension, or claim national representativeness.')
page_break()

heading('CHAPTER 2. LITERATURE REVIEW AND CONCEPTUAL BASIS',1)
heading('2.1 BMI and screening interpretation',2)
p('BMI is calculated from body weight divided by squared height. It is useful for population screening but does not directly measure abdominal adiposity, body composition, medication use, diet, or physical activity. Therefore, it should be interpreted as one screening characteristic rather than a complete cardiometabolic profile.')
equation('Equation 1. Body mass index.', 'BMI = weight (kg) / [height (m)]²')
heading('2.2 Blood pressure and the screening flag',2)
p('SBP is analysed continuously because it preserves information across the observed range. The dataset also contains a binary screening flag based on SBP/DBP ≥140/90. This flag is used for descriptive and association analyses only; it is not treated as a confirmed hypertension diagnosis because clinical confirmation requires appropriate measurement and clinical assessment.')
heading('2.3 Analytical evidence in Vietnam',2)
p('Vietnamese cross-sectional studies have reported relationships between anthropometric indices and elevated blood pressure. Nguyen Minh and Nguyen Vo examined anthropometric indices in Vietnamese adults; Wang and colleagues highlighted the relevance of BMI and waist circumference; Nguyen Duc and colleagues reported BMI- and waist-related blood-pressure patterns in northern mountainous populations. These studies provide context, but their settings, measurement protocols, covariates, and samples differ from the present dataset.')
heading('2.4 Conceptual model',2)
p('The analytical model treats SBP as the outcome, BMI as the principal explanatory variable, and age, sex, and pulse as prespecified covariates. DBP is intentionally not entered as a predictor because it is a concurrent blood-pressure measurement rather than a background covariate for the stated SBP question.')
page_break()

heading('CHAPTER 3. DATA AND METHODS',1)
heading('3.1 Study design and data governance',2)
p('This is a cross-sectional secondary analysis of de-identified community periodic health-screening records collected in 2026. The final analytical dataset contains 7,146 encounters and 14 variables. Direct identifiers were excluded before analysis. A research identifier was retained only to audit uniqueness; source-repeat status was retained as a quality flag for sensitivity analysis.')
heading('3.2 Sample construction',2)
figure('05_cohort_flow.png','Figure 3.1. Transparent sample construction from non-empty source records to the final analytical sample.',5.7)
p('Records were progressively restricted to adults aged 18–59 years with required basic measurements. The flow diagram documents the selection process: 10,912 non-empty source records, 10,756 records aged 18–59 years, 7,174 records with core measurements, and 7,146 records in the final analytical sample.')
heading('3.3 Variables and operational definitions',2)
table(['Role','Variable','Operational use'],[
 ['Outcome','SBP (mmHg)','Continuous outcome in regression and group comparisons'],
 ['Main exposure','BMI (kg/m²)','Continuous and five-level analytical group'],
 ['Covariates','Age, sex, pulse','Prespecified adjustment variables'],
 ['Secondary outcome','Screening flag ≥140/90','Binary descriptive/association indicator, not diagnosis'],
 ['Quality flag','Source repeat flag','Sensitivity-analysis exclusion indicator']])
heading('3.4 Data quality and outlier audit',2)
p('The workflow checked record count, variable count, data types, missingness, research-ID duplication, and the source-repeat flag. Numerical outliers were flagged using the conventional IQR rule. Flagging does not imply erroneous measurement; values were retained unless there was evidence of data error.')
equation('Equation 2. IQR outlier limits.', 'Lower = Q1 − 1.5 × IQR;  Upper = Q3 + 1.5 × IQR')
heading('3.5 Statistical analysis',2)
p('BMI and SBP were assessed using histogram–kernel density plots, Q–Q plots, skewness, excess kurtosis, D’Agostino–Pearson normality tests, and Shapiro–Wilk tests on a reproducible sample of up to 5,000 encounters. Given the large sample, p-values were interpreted together with graphical and descriptive evidence.')
p('Welch t-tests compared SBP between sexes. Classical one-way ANOVA was reported for course coverage; Levene tests assessed variance homogeneity and Welch ANOVA plus pairwise Welch tests with Holm adjustment were used where variances differed. Chi-square testing evaluated BMI group × screening flag, after auditing expected frequencies. Pearson and Spearman statistics quantified linear and monotonic association, respectively.')
equation('Equation 3. Pearson correlation coefficient.', 'r = Σ[(xi − x̄)(yi − ȳ)] / √[Σ(xi − x̄)² × Σ(yi − ȳ)²]')
heading('3.6 Multiple linear regression and diagnostics',2)
equation('Equation 4. Main regression model.', 'SBPᵢ = β₀ + β₁BMIᵢ + β₂Ageᵢ + β₃Maleᵢ + β₄Pulseᵢ + εᵢ')
p('The model reports coefficients, standard errors, t statistics, p-values, 95% confidence intervals, R², adjusted R², RMSE, VIF, residual-versus-fitted plots, and a residual Q–Q plot. A sensitivity analysis re-estimated the BMI coefficient after excluding encounters with the source-repeat flag.')
page_break()

heading('CHAPTER 4. RESULTS',1)
heading('4.1 Data quality and descriptive overview',2)
p('The final analytical file contains 7,146 encounters and 14 variables. No missing values were present in the released analytical variables, and research identifiers were unique. There were 118 source-repeat-flagged encounters. These were retained in the principal analysis and separately evaluated in the sensitivity analysis.')
desc=pd.read_csv(TAB/'01_descriptive_statistics.csv',index_col=0)
show=[]
for var,label in [('age','Age (years)'),('bmi','BMI (kg/m²)'),('pulse','Pulse (beats/min)'),('sbp','SBP (mmHg)'),('dbp','DBP (mmHg)')]:
 r=desc.loc[var]; show.append([label, f"{r['mean']:.2f}",f"{r['median']:.2f}",f"{r['std']:.2f}",f"{r['min']:.0f}–{r['max']:.0f}"])
table(['Variable','Mean','Median','SD','Range'],show)
figure('06_sample_composition.png','Figure 4.1. Sample composition by sex and age group.',5.9)
heading('4.2 Distribution assessment',2)
figure('01_distributions_bmi_sbp.png','Figure 4.2. Histogram and kernel density estimates for BMI and SBP.',5.9)
dist=pd.read_csv(TAB/'08_distribution_audit.csv')
table(['Variable','Skewness','Excess kurtosis','Interpretation'],[[r['Biến'],f"{r['Skewness']:.3f}",f"{r['Excess kurtosis']:.3f}",'Right-skewed; not fully normal'] for _,r in dist.iterrows()])
p('Both distributions departed from strict normality. BMI had skewness 1.058 and excess kurtosis 2.997; SBP had skewness 0.820 and excess kurtosis 1.284. Normality tests were statistically significant, but their sensitivity at large n was expected. Q–Q plots confirmed tail deviation, motivating robust complementary analyses.')
figure('13_distribution_diagnostics.png','Figure 4.3. Q–Q diagnostics for BMI and SBP.',5.9)
heading('4.3 Group differences and screening-pattern association',2)
p('SBP differed by sex in the Welch t-test (t = 26.581, p < 0.001; Cohen’s d = 0.629). Classical ANOVA found SBP differences across BMI groups (F = 162.401, p < 0.001; η² = 0.083) and age groups (F = 186.611, p < 0.001; η² = 0.073). Levene tests indicated unequal variances, and Welch ANOVA confirmed the group-level conclusions.')
tests=pd.read_csv(TAB/'08_hypothesis_tests_summary.csv')
table(['Analysis','Test statistic','p-value','Effect size'],[[r['Phân tích'],f"{r['Thống kê']:.3f}",'<0.001' if r['p-value']<.001 else f"{r['p-value']:.3f}",f"{r['Kích thước hiệu ứng']}: {r['Giá trị hiệu ứng']:.3f}"] for _,r in tests.iterrows()])
figure('02_bmi_groups_and_scatter.png','Figure 4.4. SBP by BMI group and the continuous BMI–SBP pattern.',5.9)
figure('03_screening_rate_by_bmi.png','Figure 4.5. Screening-flag prevalence increased across BMI groups.',5.5)
p('The screening flag was associated with BMI group (χ² = 337.862, p < 0.001; Cramér’s V = 0.217). The flag proportion rose from 6.6% in BMI <18.5 to 38.3% in BMI ≥30. Expected cell frequencies met the Chi-square requirement, with a minimum expected frequency above 5.')
figure('09_screening_heatmap_age_bmi.png','Figure 4.6. Screening-flag prevalence by age group and BMI group.',5.8)
heading('4.4 Correlation analysis',2)
p('BMI and SBP showed a positive association: Pearson r = 0.287 and Spearman ρ = 0.335, both p < 0.001. The magnitude is low-to-moderate rather than deterministic, indicating that BMI is one correlate among multiple contributors to SBP variation.')
figure('10_spearman_correlation_heatmap.png','Figure 4.7. Spearman correlation matrix of continuous analytical variables.',5.3)
heading('4.5 Adjusted regression findings',2)
reg=pd.read_csv(TAB/'06_adjusted_ols_regression.csv')
table(['Term','β','95% CI','Interpretation'],[[r['Biến'],f"{r['Hệ số β']:.3f}",f"{r['CI 95% thấp']:.3f} to {r['CI 95% cao']:.3f}",'p < 0.001'] for _,r in reg.iterrows()])
p('The adjusted model had R² = 0.257, adjusted R² = 0.256, and RMSE = 14.07 mmHg. Holding age, sex, and pulse constant, a one-unit higher BMI was associated with 0.861 mmHg higher SBP. The model explains a meaningful but incomplete proportion of variation; unmeasured factors such as medication, diet, smoking, alcohol use, physical activity, and comorbidity remain possible sources of residual variation.')
figure('11_regression_forest_plot.png','Figure 4.8. Adjusted SBP differences expressed in practical increments with 95% confidence intervals.',5.6)
figure('04_regression_diagnostics.png','Figure 4.9. Regression residual diagnostics.',5.9)
heading('4.6 Sensitivity analysis',2)
p('The BMI coefficient was 0.861 mmHg/1 BMI in the main analysis and 0.852 after excluding the 118 repeat-flagged encounters. This small change supports robustness of the principal association to this data-quality decision.')
figure('12_sensitivity_repeat_flag.png','Figure 4.10. Stability of the BMI coefficient in the repeat-flag sensitivity analysis.',5.4)
page_break()

heading('CHAPTER 5. DISCUSSION',1)
heading('5.1 Principal findings',2)
p('This analysis found a consistent positive statistical relationship between BMI and SBP in de-identified community screening encounters. The association appeared in continuous correlation analysis, BMI-group comparisons, the screening-flag gradient, and adjusted regression. The sensitivity analysis produced a nearly unchanged BMI coefficient after excluding repeat-flagged encounters.')
heading('5.2 Interpretation in context',2)
p('The observed direction is consistent with Vietnamese cross-sectional studies linking anthropometric measures to blood-pressure outcomes. Nevertheless, the present analysis should not be treated as a direct replication because its population, screening pathway, covariate availability, and data-cleaning rules differ. In particular, absence of waist circumference prevents a comparison between BMI and central-adiposity indicators.')
heading('5.3 Practical significance',2)
p('The pattern may assist programme-level prioritisation. For example, the age-by-BMI heatmap indicates higher screening-flag prevalence in several combinations of older age and higher BMI. This can guide outreach or follow-up planning, but it must not be converted into an automated diagnosis, a treatment rule, or advice for an individual without clinical assessment.')
heading('5.4 Strengths',2)
bullets(['A large analytical sample with one data source used consistently across all rubric components.', 'Explicit audits of completeness, uniqueness, outliers, expected frequencies, variance homogeneity and residual patterns.', 'Reporting of confidence intervals and effect sizes alongside p-values.', 'Welch ANOVA and Holm-adjusted follow-up tests added when equal-variance assumptions were not supported.', 'Sensitivity analysis and visual stratification added beyond minimum course requirements.'])
heading('5.5 Limitations',2)
bullets(['The cross-sectional design cannot establish that BMI causes an SBP change.', 'An encounter-level dataset may not fully capture person-level longitudinal dependence.', 'The screening flag does not confirm hypertension.', 'Lifestyle, medication, waist circumference and clinical-comorbidity data were unavailable.', 'The screening sample is not a probability sample of all Vietnamese adults.', 'Residual Q–Q tails depart from normality; regression inference should be interpreted with appropriate caution.'])
page_break()

heading('CHAPTER 6. CONCLUSION AND RECOMMENDATIONS',1)
heading('6.1 Conclusion',2)
p('The project meets the full analytical requirements of exploratory analysis, probability-distribution assessment, hypothesis testing, correlation analysis and multiple linear regression using one consistent dataset. BMI was positively associated with SBP after adjustment for age, sex and pulse. The screening-flag proportion rose substantially across BMI groups, and the principal BMI association remained stable in the repeat-flag sensitivity analysis.')
heading('6.2 Recommendations',2)
bullets(['Future data collection should add waist circumference, medication, smoking, alcohol use, dietary factors, physical activity and comorbidity variables.', 'A future longitudinal design should distinguish repeated persons and assess within-person change over time.', 'Screening programmes may use the aggregate patterns to prioritise education and follow-up, while preserving clinician judgement and consent boundaries.', 'Any external dissemination should remain de-identified and follow the access controls applied to the research dataset.'])
heading('6.3 Reproducibility statement',2)
p('The submitted notebook contains executable Python code, figures and aggregate result tables. It is configured to read a de-identified CSV through data/Research_Data.csv or an M3_DATA_PATH environment variable. The identifiable source workbook is not included. The analysis is reproducible for authorised users with access to the de-identified research CSV.')
page_break()

heading('REFERENCES',1)
for ref in [
'Nguyen Minh, Q., & Nguyen Vo, M. (2020). Anthropometric indexes for predicting high blood pressure in Vietnamese adults: A cross-sectional study. Integrative Blood Pressure Control, 13, 181–186. https://doi.org/10.2147/IBPC.S281996.',
'Wang, J., et al. (2018). The importance of waist circumference and body mass index in cross-sectional relationships with risk of cardiovascular disease in Vietnam. PLoS ONE. PMID: 29813112.',
'Nguyen Duc, N., et al. (2019). Impact of body mass index and waist circumference on blood pressure: A cross-sectional survey in a population living in the Vietnam northern mountainous. Diabetes & Metabolic Syndrome, 13(2), 1399–1404. https://doi.org/10.1016/j.dsx.2019.02.015.',
'American Statistical Association. (2016). ASA statement on statistical significance and p-values. The American Statistician, 70(2), 129–133.'
]: p(ref,after=8)
page_break()

heading('APPENDICES',1)
heading('Appendix A. Data dictionary',2)
table(['Vietnamese field','English analytical name','Unit/format','Role'],[
['Tuổi (năm)','age','years','Covariate'],['Giới tính','sex','categorical','Covariate'],['Chiều cao (cm)','height_cm','cm','Descriptive'],['Cân nặng (kg)','weight_kg','kg','Descriptive'],['BMI (kg/m²)','bmi','kg/m²','Main exposure'],['Mạch (lần/phút)','pulse','beats/min','Covariate'],['Huyết áp tâm thu (mmHg)','sbp','mmHg','Primary outcome'],['Huyết áp tâm trương (mmHg)','dbp','mmHg','Descriptive/screening input'],['Cờ sàng lọc HA ≥140/90','bp_screen','0/1','Secondary outcome'],['Cờ mã nguồn lặp','repeat_flag','0/1','Sensitivity flag']])
heading('Appendix B. Additional figures',2)
for f,cap,w in [
('07_bmi_by_sex.png','Figure B1. BMI distribution by sex.',5.2),
('08_sbp_by_age_and_sex.png','Figure B2. Mean SBP by age group and sex.',5.8)]: figure(f,cap,w)
heading('Appendix C. Additional statistical tables',2)
ass=pd.read_csv(TAB/'09_assumption_audit.csv')
table(list(ass.columns), ass.fillna('').values.tolist(), size=8.5)
post=pd.read_csv(TAB/'10_posthoc_bmi_welch_holm.csv')
table(list(post.columns), post.round(4).values.tolist(), size=7.5)

doc.save(DEST)
print(DEST)
